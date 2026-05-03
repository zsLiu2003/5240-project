from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


BASE = Path(__file__).resolve().parent
PPTX = BASE / "5240_pre_Group_9_260502.pptx"
BACKUP = BASE / "5240_pre_Group_9_260502.backup.pptx"

BLUE = RGBColor(0, 55, 110)
MID_BLUE = RGBColor(0, 91, 170)
LIGHT_BLUE = RGBColor(235, 246, 255)
YELLOW = RGBColor(245, 176, 0)
GREEN = RGBColor(28, 112, 65)
LIGHT_GREEN = RGBColor(238, 249, 241)
PURPLE = RGBColor(64, 43, 140)
LIGHT_PURPLE = RGBColor(244, 242, 255)
ORANGE = RGBColor(206, 139, 0)
LIGHT_ORANGE = RGBColor(255, 248, 235)
RED = RGBColor(179, 38, 38)
LIGHT_RED = RGBColor(255, 240, 240)
BLACK = RGBColor(0, 0, 0)
GRAY = RGBColor(110, 110, 110)
WHITE = RGBColor(255, 255, 255)


def inch(value):
    return Inches(value)


def set_text(shape, text, size=20, color=BLACK, bold=False, align=None):
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = inch(0.08)
    tf.margin_right = inch(0.08)
    tf.margin_top = inch(0.03)
    tf.margin_bottom = inch(0.03)
    p = tf.paragraphs[0]
    if align is not None:
        p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def add_text(slide, x, y, w, h, text, size=20, color=BLACK, bold=False, align=None):
    shape = slide.shapes.add_textbox(inch(x), inch(y), inch(w), inch(h))
    return set_text(shape, text, size, color, bold, align)


def add_box(slide, x, y, w, h, text="", fill=WHITE, line=BLUE, radius=True, width=1.2):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    box = slide.shapes.add_shape(shape_type, inch(x), inch(y), inch(w), inch(h))
    box.fill.solid()
    box.fill.fore_color.rgb = fill
    box.line.color.rgb = line
    box.line.width = Pt(width)
    if text:
        set_text(box, text)
    return box


def add_line(slide, x1, y1, x2, y2, color=BLACK, width=2.0, arrow=False):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, inch(x1), inch(y1), inch(x2), inch(y2)
    )
    line.line.color.rgb = color
    line.line.width = Pt(width)
    if arrow:
        line.line.end_arrowhead = True
    return line


def add_header(slide, title, page):
    # Top blue rule and HKUST-like wordmark are native shapes/text.
    top = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, inch(13.333), inch(0.22))
    top.fill.solid()
    top.fill.fore_color.rgb = BLUE
    top.line.fill.background()
    mark = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, inch(0.36), inch(0.36), inch(0.68))
    mark.fill.solid()
    mark.fill.fore_color.rgb = YELLOW
    mark.line.fill.background()
    add_text(slide, 0.55, 0.48, 8.8, 0.55, title, 30, BLACK, True)
    add_text(slide, 10.35, 0.38, 2.45, 0.56, "THE HONG KONG\nUNIVERSITY OF SCIENCE\nAND TECHNOLOGY", 13, BLUE, False)
    # Minimal editable emblem.
    emblem_dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, inch(10.05), inch(0.38), inch(0.18), inch(0.18))
    emblem_dot.fill.solid()
    emblem_dot.fill.fore_color.rgb = YELLOW
    emblem_dot.line.fill.background()
    add_text(slide, 12.88, 7.15, 0.28, 0.18, str(page), 12, GRAY)


def add_bottom_note(slide, x, y, w, text, icon="bulb"):
    box = add_box(slide, x, y, w, 0.68, "", WHITE, MID_BLUE, True, 1.0)
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, inch(x + 0.35), inch(y + 0.12), inch(0.45), inch(0.45))
    circle.fill.solid()
    circle.fill.fore_color.rgb = BLUE
    circle.line.fill.background()
    if icon == "target":
        add_target(slide, x + 0.58, y + 0.34, 0.19, WHITE)
    else:
        add_bulb(slide, x + 0.58, y + 0.34, 0.19, WHITE)
    add_text(slide, x + 1.0, y + 0.19, w - 1.25, 0.32, text, 19, BLACK, True)
    return box


def add_bulb(slide, cx, cy, r, color=BLUE):
    bulb = slide.shapes.add_shape(MSO_SHAPE.OVAL, inch(cx - r * 0.62), inch(cy - r * 0.75), inch(r * 1.24), inch(r * 1.12))
    bulb.fill.solid()
    bulb.fill.fore_color.rgb = color
    bulb.line.fill.background()
    add_line(slide, cx - r * 0.35, cy + r * 0.1, cx + r * 0.35, cy + r * 0.1, color, 1.5)
    add_line(slide, cx - r * 0.25, cy + r * 0.28, cx + r * 0.25, cy + r * 0.28, color, 1.5)


def add_target(slide, cx, cy, r, color=BLUE):
    for k in (1.0, 0.62, 0.28):
        c = slide.shapes.add_shape(MSO_SHAPE.OVAL, inch(cx - r * k), inch(cy - r * k), inch(2 * r * k), inch(2 * r * k))
        c.fill.background()
        c.line.color.rgb = color
        c.line.width = Pt(2)
    add_line(slide, cx, cy, cx + r * 1.25, cy - r * 1.25, color, 2)


def add_dot(slide, x, y, size=0.1, color=BLUE):
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, inch(x), inch(y), inch(size), inch(size))
    dot.fill.solid()
    dot.fill.fore_color.rgb = color
    dot.line.fill.background()
    return dot


def add_small_molecule(slide, cx, cy, scale=1.0, color=MID_BLUE):
    pts = [(cx, cy), (cx - 0.25 * scale, cy + 0.18 * scale), (cx + 0.26 * scale, cy + 0.2 * scale), (cx + 0.07 * scale, cy - 0.28 * scale)]
    for x1, y1 in pts[1:]:
        add_line(slide, cx, cy, x1, y1, color, 1.5)
    for x, y in pts:
        c = slide.shapes.add_shape(MSO_SHAPE.OVAL, inch(x - 0.055 * scale), inch(y - 0.055 * scale), inch(0.11 * scale), inch(0.11 * scale))
        c.fill.solid()
        c.fill.fore_color.rgb = WHITE if (x, y) == (cx, cy) else RGBColor(64, 145, 220)
        c.line.color.rgb = color
        c.line.width = Pt(1.2)


def add_hexagon(slide, cx, cy, r=0.24, color=GREEN):
    pts = [
        (cx + r, cy),
        (cx + r * 0.5, cy + r * 0.86),
        (cx - r * 0.5, cy + r * 0.86),
        (cx - r, cy),
        (cx - r * 0.5, cy - r * 0.86),
        (cx + r * 0.5, cy - r * 0.86),
    ]
    for i in range(6):
        add_line(slide, pts[i][0], pts[i][1], pts[(i + 1) % 6][0], pts[(i + 1) % 6][1], color, 1.8)
    for x, y in pts:
        c = slide.shapes.add_shape(MSO_SHAPE.OVAL, inch(x - 0.045), inch(y - 0.045), inch(0.09), inch(0.09))
        c.fill.solid()
        c.fill.fore_color.rgb = LIGHT_GREEN
        c.line.color.rgb = color


def add_bars(slide, x, y, color=ORANGE):
    for i, h in enumerate((0.16, 0.28, 0.42)):
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, inch(x + i * 0.13), inch(y - h), inch(0.08), inch(h))
        bar.fill.solid()
        bar.fill.fore_color.rgb = color
        bar.line.fill.background()
    add_line(slide, x - 0.03, y, x + 0.43, y, color, 1.6)


def add_lock(slide, x, y, locked=True, color=BLUE):
    body = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, inch(x), inch(y + 0.11), inch(0.26), inch(0.22))
    body.fill.solid()
    body.fill.fore_color.rgb = color
    body.line.fill.background()
    if locked:
        add_line(slide, x + 0.06, y + 0.13, x + 0.06, y + 0.02, color, 2)
        add_line(slide, x + 0.06, y + 0.02, x + 0.20, y + 0.02, color, 2)
        add_line(slide, x + 0.20, y + 0.02, x + 0.20, y + 0.13, color, 2)
    else:
        add_line(slide, x + 0.06, y + 0.13, x + 0.06, y + 0.02, color, 2)
        add_line(slide, x + 0.06, y + 0.02, x + 0.23, y - 0.03, color, 2)


def draw_architecture(slide):
    add_header(slide, "ChemBERTa Regression Architecture", 9)
    panel = add_box(slide, 0.82, 1.38, 4.55, 4.48, "", WHITE, MID_BLUE, True)
    head = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, inch(0.82), inch(1.38), inch(4.55), inch(0.62))
    head.fill.solid()
    head.fill.fore_color.rgb = BLUE
    head.line.color.rgb = BLUE
    set_text(head, "Key Setup", 22, WHITE, True)
    items = [
        ("Input:", "SMILES string"),
        ("Target:", "Energy_min binding energy"),
        ("Model:", "ChemBERTa encoder +\nregression head"),
        ("Task type:", "supervised regression"),
    ]
    for i, (k, v) in enumerate(items):
        y = 2.42 + i * 0.76
        add_dot(slide, 1.05, y + 0.08, 0.09, BLUE)
        if i == 0:
            add_small_molecule(slide, 1.52, y + 0.12, 1.1, BLUE)
        elif i == 1:
            add_target(slide, 1.53, y + 0.16, 0.24, BLUE)
        elif i == 2:
            add_hexagon(slide, 1.53, y + 0.16, 0.20, BLUE)
        else:
            add_line(slide, 1.25, y + 0.42, 1.82, y + 0.42, BLUE, 1.8)
            add_line(slide, 1.25, y + 0.42, 1.25, y - 0.03, BLUE, 1.8)
            add_line(slide, 1.35, y + 0.34, 1.55, y + 0.18, BLUE, 1.8)
            add_line(slide, 1.55, y + 0.18, 1.77, y + 0.02, BLUE, 1.8)
        add_text(slide, 2.02, y - 0.02, 0.78, 0.32, k, 17, BLUE, True)
        add_text(slide, 2.8, y - 0.02, 2.1, 0.46, v, 17, BLACK)

    steps = [
        ("SMILES", LIGHT_BLUE, MID_BLUE),
        ("Tokenizer", LIGHT_BLUE, RGBColor(40, 150, 210)),
        ("ChemBERTa Encoder", LIGHT_GREEN, GREEN),
        ("[CLS] Molecular Representation", LIGHT_PURPLE, PURPLE),
        ("Regression Head", LIGHT_ORANGE, ORANGE),
        ("Predicted Binding Energy", LIGHT_RED, RED),
    ]
    x, w, h = 6.35, 5.85, 0.58
    for i, (txt, fill, line) in enumerate(steps):
        y = 1.38 + i * 0.76
        add_box(slide, x, y, w, h, "", fill, line, True, 1.0)
        if i == 0:
            add_hexagon(slide, x + 0.7, y + 0.29, 0.16, MID_BLUE)
        elif i == 1:
            for j, tok in enumerate(["C", "1", "=", "C", "C", "N", "..."]):
                t = add_box(slide, x + 0.18 + j * 0.29, y + 0.19, 0.24, 0.24, tok, WHITE, RGBColor(35, 140, 215), True)
                set_text(t, tok, 12, MID_BLUE, True, PP_ALIGN.CENTER)
        elif i == 2:
            add_hexagon(slide, x + 0.55, y + 0.29, 0.19, GREEN)
        elif i == 3:
            tag = add_box(slide, x + 0.2, y + 0.18, 0.72, 0.28, "[CLS]", WHITE, PURPLE, True)
            set_text(tag, "[CLS]", 15, PURPLE, True, PP_ALIGN.CENTER)
        elif i == 4:
            add_bars(slide, x + 0.3, y + 0.42, ORANGE)
        else:
            add_target(slide, x + 0.52, y + 0.3, 0.22, RED)
        add_text(slide, x + 2.1, y + 0.15, 3.1, 0.3, txt, 18, BLACK, True, PP_ALIGN.CENTER)
        if i < len(steps) - 1:
            add_line(slide, x + w / 2, y + h, x + w / 2, y + 0.76, BLACK, 1.8, True)
    add_bottom_note(slide, 1.62, 6.18, 10.0, "ChemBERTa provides pretrained molecular SMILES representations.")


def draw_smiles_aug(slide):
    add_header(slide, "SMILES Augmentation: Invariance to String Serialization", 10)
    left = add_box(slide, 0.36, 1.46, 4.16, 4.1, "", WHITE, MID_BLUE, True)
    add_text(slide, 0.67, 1.67, 2.0, 0.34, "Key Points", 22, BLUE, True)
    add_line(slide, 0.52, 2.08, 4.38, 2.08, MID_BLUE, 1.2)
    points = [
        "One molecule can have multiple\nvalid SMILES strings.",
        "During training: canonical SMILES ->\nrandomized equivalent SMILES",
        "The binding-energy label\nstays unchanged.",
        "Goal: reduce sensitivity to\narbitrary SMILES ordering",
    ]
    for i, txt in enumerate(points):
        y = 2.45 + i * 0.82
        if i == 0:
            add_small_molecule(slide, 0.9, y + 0.14, 1.1, BLUE)
        elif i == 1:
            add_line(slide, 0.62, y + 0.16, 1.05, y + 0.16, BLUE, 2.0, True)
            add_line(slide, 1.05, y + 0.28, 0.62, y + 0.28, BLUE, 2.0, True)
        elif i == 2:
            tag = slide.shapes.add_shape(MSO_SHAPE.PENTAGON, inch(0.62), inch(y), inch(0.45), inch(0.35))
            tag.fill.solid(); tag.fill.fore_color.rgb = BLUE; tag.line.fill.background()
        else:
            add_target(slide, 0.86, y + 0.16, 0.22, BLUE)
        add_text(slide, 1.38, y - 0.06, 2.85, 0.45, txt, 15.5, BLACK)

    variants = [
        ("SMILES variant 1", "C1=CC=CC=N1", MID_BLUE, LIGHT_BLUE),
        ("SMILES variant 2", "N1=CC=CC=C1", GREEN, LIGHT_GREEN),
        ("SMILES variant 3", "C=1C=NC=CC=1", PURPLE, LIGHT_PURPLE),
    ]
    for i, (label, code, color, fill) in enumerate(variants):
        y = 1.55 + i * 1.33
        add_box(slide, 4.95, y, 3.25, 1.02, "", fill, color, True)
        add_text(slide, 5.18, y + 0.18, 1.8, 0.24, label, 16, color, True)
        add_text(slide, 5.18, y + 0.55, 1.85, 0.28, code, 17, BLACK)
        add_hexagon(slide, 7.58, y + 0.52, 0.31, color)
        add_line(slide, 8.2, y + 0.51, 9.25, 3.37, RGBColor(36, 48, 70), 2.2, True)
    mol = slide.shapes.add_shape(MSO_SHAPE.OVAL, inch(9.08), inch(2.58), inch(1.52), inch(1.52))
    mol.fill.solid(); mol.fill.fore_color.rgb = RGBColor(242, 247, 255); mol.line.color.rgb = MID_BLUE
    add_small_molecule(slide, 9.84, 3.05, 1.0, MID_BLUE)
    add_text(slide, 9.52, 3.48, 0.62, 0.4, "same\nmolecule", 13, BLACK, True, PP_ALIGN.CENTER)
    add_line(slide, 10.6, 3.34, 11.3, 3.34, RGBColor(36, 48, 70), 2.2, True)
    out = add_box(slide, 11.38, 2.7, 1.55, 1.25, "", LIGHT_ORANGE, ORANGE, True)
    add_bars(slide, 11.96, 3.16, ORANGE)
    add_text(slide, 11.55, 3.48, 1.18, 0.38, "same\nEnergy_min label", 13, BLACK, True, PP_ALIGN.CENTER)
    add_bottom_note(slide, 1.52, 5.92, 9.95, "Augmentation changes the input view, not the supervised target.")


def draw_two_stage(slide):
    add_header(slide, "Two-Stage Supervised Fine-Tuning", 11)
    add_box(slide, 0.42, 1.30, 8.78, 1.9, "", WHITE, MID_BLUE, True)
    add_box(slide, 0.42, 3.58, 8.78, 1.85, "", WHITE, GREEN, True)
    add_text(slide, 0.68, 1.62, 2.8, 0.3, "Stage 1: head warm-up", 20, BLUE, True)
    add_text(slide, 0.68, 2.10, 3.05, 0.74, "Freeze the ChemBERTa\nencoder and train only the\nregression head.", 17, BLACK)
    add_text(slide, 0.68, 3.92, 3.0, 0.3, "Stage 2: task adaptation", 20, GREEN, True)
    add_text(slide, 0.68, 4.40, 3.05, 0.74, "Unfreeze the ChemBERTa\nencoder and fine-tune\nencoder plus regression head.", 17, BLACK)

    for y, locked, line in [(1.55, True, GREEN), (3.82, False, GREEN)]:
        enc = add_box(slide, 4.12, y, 2.48, 1.36, "", LIGHT_GREEN, line, True)
        add_hexagon(slide, 4.72, y + 0.47, 0.24, line)
        add_text(slide, 4.96, y + 0.72, 1.2, 0.38, "ChemBERTa\nEncoder", 16, BLACK, True, PP_ALIGN.CENTER)
        add_lock(slide, 6.22, y + 0.14, locked, BLUE if locked else GREEN)
        add_line(slide, 6.6, y + 0.68, 7.1, y + 0.68, BLACK, 1.8, True)
        head = add_box(slide, 7.1, y, 1.85, 1.36, "", LIGHT_ORANGE, ORANGE, True)
        add_bars(slide, 7.72, y + 0.62, ORANGE)
        add_text(slide, 7.48, y + 0.77, 1.1, 0.36, "Regression\nHead", 15, BLACK, True, PP_ALIGN.CENTER)
        add_text(slide, 8.45, y + 0.14, 0.3, 0.3, "LR", 12, ORANGE, True)
    add_line(slide, 5.05, 3.2, 5.05, 3.58, BLUE, 4.0, True)

    lr = add_box(slide, 9.75, 1.98, 3.18, 2.95, "", WHITE, MID_BLUE, True)
    add_bulb(slide, 10.42, 2.54, 0.25, BLUE)
    add_text(slide, 10.86, 2.38, 1.6, 0.3, "Learning rates", 18, BLUE, True)
    add_line(slide, 9.98, 2.98, 12.7, 2.98, MID_BLUE, 1.0)
    for i, txt in enumerate(["smaller LR for ChemBERTa\nbackbone", "larger LR for regression head"]):
        y = 3.48 + i * 0.8
        add_dot(slide, 10.0, y, 0.1, BLUE)
        add_text(slide, 10.32, y - 0.05, 2.0, 0.42, txt, 14.5, BLACK)
    add_bottom_note(slide, 1.3, 5.7, 10.1, "Stage 1 learns the supervised energy scale; Stage 2 adapts the molecular representation.", "target")


def draw_experimental_design(slide):
    add_header(slide, "Controlled Experimental Design", 12)
    table_x, table_y, table_w = 0.36, 1.6, 7.8
    row_h = 0.78
    col1, col2 = 3.15, 4.65
    add_box(slide, table_x, table_y, table_w, row_h, "", BLUE, BLUE, False)
    add_text(slide, table_x + 1.25, table_y + 0.22, 1.1, 0.28, "Method", 18, WHITE, True, PP_ALIGN.CENTER)
    add_text(slide, table_x + col1 + 1.55, table_y + 0.22, 1.0, 0.28, "Purpose", 18, WHITE, True, PP_ALIGN.CENTER)
    rows = [
        ("No Aug + Two-stage", "Baseline without SMILES augmentation"),
        ("Aug + Two-stage", "Final method"),
        ("Aug + Stage1-only", "Tests whether frozen ChemBERTa\nfeatures are sufficient"),
        ("Aug + Stage2-only", "Tests whether Stage 1 warm-up is useful"),
    ]
    for i, (method, purpose) in enumerate(rows):
        y = table_y + row_h * (i + 1)
        fill = RGBColor(245, 248, 251) if i % 2 else WHITE
        add_box(slide, table_x, y, table_w, row_h, "", fill, RGBColor(190, 205, 220), False, 0.7)
        add_line(slide, table_x + col1, y, table_x + col1, y + row_h, RGBColor(190, 205, 220), 0.8)
        add_text(slide, table_x + 0.55, y + 0.22, 2.2, 0.28, method, 16, BLACK, True, PP_ALIGN.CENTER)
        add_text(slide, table_x + col1 + 0.45, y + 0.18, 3.7, 0.38, purpose, 16, BLACK)

    right = add_box(slide, 8.42, 1.65, 4.58, 3.65, "", WHITE, BLUE, True)
    bullets = [
        "3 seeds x 5 folds = 15 runs per method",
        "No Aug vs Aug: effect of SMILES\naugmentation",
        "Stage1-only vs Two-stage: need for\nencoder adaptation",
        "Stage2-only vs Two-stage: role of\nStage 1 warm-up",
    ]
    for i, txt in enumerate(bullets):
        y = 2.15 + i * 0.82
        add_dot(slide, 8.68, y, 0.1, BLUE)
        add_text(slide, 9.02, y - 0.08, 3.3, 0.48, txt, 15.5, BLACK, True if i == 0 else False)
    add_bottom_note(slide, 1.95, 5.63, 9.45, "These controlled comparisons test why each training component matters.")
    bottom = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, inch(7.13), inch(13.333), inch(0.37))
    bottom.fill.solid(); bottom.fill.fore_color.rgb = BLUE; bottom.line.fill.background()
    add_text(slide, 12.88, 7.18, 0.28, 0.18, "12", 12, WHITE)


def delete_slide(prs, index):
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    r_id = slides[index].rId
    prs.part.drop_rel(r_id)
    xml_slides.remove(slides[index])


def renumber(prs):
    for i, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if not hasattr(shape, "text"):
                continue
            text = shape.text.strip()
            if text.isdigit() and len(text) <= 2:
                shape.text = str(i)
                try:
                    for para in shape.text_frame.paragraphs:
                        for run in para.runs:
                            run.font.name = "Arial"
                            run.font.size = Pt(12)
                except Exception:
                    pass


def main():
    if not BACKUP.exists():
        raise FileNotFoundError(f"Missing backup: {BACKUP}")
    PPTX.write_bytes(BACKUP.read_bytes())
    prs = Presentation(str(PPTX))
    blank = prs.slide_layouts[6]
    makers = [draw_architecture, draw_smiles_aug, draw_two_stage, draw_experimental_design]

    new_ids = []
    for maker in makers:
        slide = prs.slides.add_slide(blank)
        maker(slide)
        new_ids.append(prs.slides._sldIdLst[-1])

    sld_id_lst = prs.slides._sldIdLst
    for slide_id in new_ids:
        sld_id_lst.remove(slide_id)
    for offset, slide_id in enumerate(new_ids):
        sld_id_lst.insert(8 + offset, slide_id)

    for _ in range(6):
        delete_slide(prs, 12)
    renumber(prs)
    prs.save(str(PPTX))


if __name__ == "__main__":
    main()
